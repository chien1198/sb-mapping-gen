#!/usr/bin/env python3
"""Extract raw design documents in input/ into per-table Markdown files in extract/.

See ../SKILL.md for the full behavior contract this script implements.
"""

import json
import re
import sys
from pathlib import Path

import docx
import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[4]
INPUT_DIR = REPO_ROOT / "input"
EXTRACT_DIR = REPO_ROOT / "extract"

XLSX_ROUTES = [
    (re.compile(r"^DATAMODEL_DWH_LOS_.*\.xlsx$"), "SB_DWH"),
    (re.compile(r"^DATAMODEL_DTM_PDTD_.*\.xlsx$"), "PDTD_DTM"),
    (re.compile(r"^Reports_.*\.xlsx$"), "Report"),
]

DOCX_ROUTE = re.compile(r"^Design_Database_.*\.docx$")


def slugify(name):
    return name.strip()


def clear_stale_markdown(out_dir: Path, keep_names):
    if not out_dir.exists():
        return
    keep_files = {f"{slugify(n)}.md" for n in keep_names}
    for f in out_dir.glob("*.md"):
        if f.name == "_index.json":
            continue
        if f.name not in keep_files:
            f.unlink()


# ---------------------------------------------------------------------------
# docx: Design_Database_*.docx -> extract/database/<TABLE_NAME>.md
# ---------------------------------------------------------------------------


def extract_design_docx(path: Path):
    d = docx.Document(str(path))
    body = d.element.body
    para_map = {p._p: p for p in d.paragraphs}
    table_map = {t._tbl: t for t in d.tables}

    heading_re = re.compile(r"^\s*[\d.]+\s*Bảng\s+(\S+)\s*$")

    tables = []  # (name, table)
    current_name = None
    for el in body:
        if el in para_map:
            p = para_map[el]
            if p.style.name == "Heading 3":
                m = heading_re.match(p.text.strip())
                current_name = m.group(1) if m else None
        elif el in table_map:
            t = table_map[el]
            if current_name is not None:
                tables.append((current_name, t))
                current_name = None  # each heading claims at most one table

    out_dir = EXTRACT_DIR / "database"
    out_dir.mkdir(parents=True, exist_ok=True)

    index = []
    seen_names = []
    for name, t in tables:
        rows = []
        for row in t.rows:
            rows.append([c.text.strip() for c in row.cells])
        if not rows:
            continue
        header = rows[0]
        data_rows = rows[1:]

        lines = [f"# {name}", "", f"Nguồn: docx \"{path.name}\"", ""]
        lines.append("| " + " | ".join(header) + " |")
        lines.append("| " + " | ".join(["---"] * len(header)) + " |")
        for r in data_rows:
            lines.append("| " + " | ".join(r) + " |")
        lines.append("")

        out_file = out_dir / f"{slugify(name)}.md"
        out_file.write_text("\n".join(lines), encoding="utf-8")
        seen_names.append(name)
        index.append(
            {
                "table": name,
                "file": out_file.name,
                "source": path.name,
                "columns": len(data_rows),
            }
        )

    clear_stale_markdown(out_dir, seen_names)
    (out_dir / "_index.json").write_text(
        json.dumps({"sources": [path.name], "tables": index}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"[database] {path.name}: extracted {len(index)} tables -> {out_dir}")
    return {t["table"] for t in index}


# ---------------------------------------------------------------------------
# xlsx: DIM/FCT datamodel sheets -> extract/<SB_DWH|PDTD_DTM>/<SHEET_NAME>.md
# ---------------------------------------------------------------------------


def sheet_to_rows(ws):
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(list(row))
    return rows


def find_column_header_row(rows):
    """Return index of the row whose first cell is exactly 'STT' (the real column table header)."""
    for i, row in enumerate(rows):
        if row and str(row[0]).strip() == "STT":
            return i
    return None


def trim_trailing_none(row):
    r = list(row)
    while r and (r[-1] is None or str(r[-1]).strip() == ""):
        r.pop()
    return r


def fmt_cell(v):
    if v is None:
        return ""
    return str(v).strip()


def extract_datamodel_xlsx(path: Path, out_subdir: str, known_table_names):
    wb = openpyxl.load_workbook(str(path), data_only=True)
    out_dir = EXTRACT_DIR / out_subdir
    out_dir.mkdir(parents=True, exist_ok=True)

    index = []
    seen_names = []
    for sheet_name in wb.sheetnames:
        if known_table_names is not None and sheet_name not in known_table_names:
            continue  # skip explainer/reference sheets not known as real tables
        ws = wb[sheet_name]
        rows = sheet_to_rows(ws)
        if not rows:
            continue

        header_idx = find_column_header_row(rows)
        if header_idx is None:
            continue  # not a table sheet in the expected shape

        meta_rows = rows[:header_idx]
        # meta block: first row is the sheet/table title, then blank row(s), then
        # "Field | Value" rows until the header. Collect non-empty "Field | Value" rows.
        meta = {}
        loai_bang = ""
        for r in meta_rows:
            r = trim_trailing_none(r)
            if len(r) >= 2 and r[0] and r[1]:
                key = fmt_cell(r[0])
                val = fmt_cell(r[1])
                if key not in ("", sheet_name):
                    meta[key] = val
        if fmt_cell(rows[0][0]) == sheet_name:
            pass  # first row is just the title, already excluded from meta via key check
        loai_bang = meta.get("Loại bảng", "")

        header = trim_trailing_none(rows[header_idx])
        data_rows = []
        for r in rows[header_idx + 1 :]:
            r = trim_trailing_none(r)
            if not r:
                continue
            if not fmt_cell(r[0]):
                continue
            data_rows.append(r)

        lines = [f"# {sheet_name}", "", f'Nguồn: xlsx sheet "{sheet_name}" ({path.name})', ""]
        for key in [
            "Loại bảng",
            "Mô tả",
            "Lưu gì",
            "Grain",
            "Khóa",
            "Nguồn",
            "Báo cáo sử dụng",
            "Quy tắc load",
            "Quy tắc ghi",
        ]:
            if key in meta:
                lines.append(f"- {key}: {meta[key]}")
        lines.append("")

        ncols = len(header)
        lines.append("| " + " | ".join(header) + " |")
        lines.append("| " + " | ".join(["---"] * ncols) + " |")
        for r in data_rows:
            padded = r + [""] * (ncols - len(r))
            lines.append("| " + " | ".join(fmt_cell(c) for c in padded[:ncols]) + " |")
        lines.append("")

        out_file = out_dir / f"{slugify(sheet_name)}.md"
        out_file.write_text("\n".join(lines), encoding="utf-8")
        seen_names.append(sheet_name)
        index.append(
            {
                "table": sheet_name,
                "file": out_file.name,
                "source": path.name,
                "loai_bang": loai_bang,
                "columns": len(data_rows),
            }
        )

    clear_stale_markdown(out_dir, seen_names)
    (out_dir / "_index.json").write_text(
        json.dumps({"sources": [path.name], "tables": index}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"[{out_subdir}] {path.name}: extracted {len(index)} sheets -> {out_dir}")


# ---------------------------------------------------------------------------
# xlsx: Reports_*.xlsx -> extract/Report/<SHEET_NAME>.md  (every sheet is real)
# ---------------------------------------------------------------------------


def extract_report_xlsx(path: Path):
    wb = openpyxl.load_workbook(str(path), data_only=True)
    out_dir = EXTRACT_DIR / "Report"
    out_dir.mkdir(parents=True, exist_ok=True)

    index = []
    seen_names = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = sheet_to_rows(ws)
        if len(rows) < 4:
            continue
        title = fmt_cell(rows[0][0]) if rows[0] else sheet_name
        note = fmt_cell(rows[1][0]) if len(rows) > 1 and rows[1] else ""
        header = trim_trailing_none(rows[3]) if len(rows) > 3 else []
        data_rows = []
        for r in rows[4:]:
            r = trim_trailing_none(r)
            if not r:
                continue
            data_rows.append(r)

        lines = [f"# {title}" if title else f"# {sheet_name}", ""]
        lines.append(f'Nguồn: xlsx sheet "{sheet_name}" ({path.name})')
        if note:
            lines.append("")
            lines.append(note)
        lines.append("")

        if header:
            ncols = len(header)
            lines.append("| " + " | ".join(fmt_cell(c) for c in header) + " |")
            lines.append("| " + " | ".join(["---"] * ncols) + " |")
            for r in data_rows:
                padded = r + [""] * (ncols - len(r))
                lines.append("| " + " | ".join(fmt_cell(c) for c in padded[:ncols]) + " |")
            lines.append("")

        out_file = out_dir / f"{slugify(sheet_name)}.md"
        out_file.write_text("\n".join(lines), encoding="utf-8")
        seen_names.append(sheet_name)
        index.append(
            {
                "table": sheet_name,
                "file": out_file.name,
                "source": path.name,
                "rows": len(data_rows),
            }
        )

    clear_stale_markdown(out_dir, seen_names)
    (out_dir / "_index.json").write_text(
        json.dumps({"sources": [path.name], "tables": index}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"[Report] {path.name}: extracted {len(index)} sheets -> {out_dir}")


# ---------------------------------------------------------------------------


def main():
    if not INPUT_DIR.exists():
        print(f"input dir not found: {INPUT_DIR}", file=sys.stderr)
        sys.exit(1)

    files = sorted(
        f
        for f in INPUT_DIR.iterdir()
        if f.is_file() and not f.name.startswith("~$") and f.suffix.lower() in (".docx", ".xlsx")
    )

    docx_files = [f for f in files if DOCX_ROUTE.match(f.name)]
    known_table_names = None
    for f in docx_files:
        names = extract_design_docx(f)
        known_table_names = names if known_table_names is None else (known_table_names | names)

    for f in files:
        if f in docx_files:
            continue
        routed = False
        for pattern, out_subdir in XLSX_ROUTES:
            if pattern.match(f.name):
                routed = True
                if out_subdir == "Report":
                    extract_report_xlsx(f)
                else:
                    extract_datamodel_xlsx(f, out_subdir, known_table_names)
                break
        if not routed:
            print(
                f"ERROR: no XLSX_ROUTES pattern matches {f.name} — add a route in extract_input.py",
                file=sys.stderr,
            )
            sys.exit(2)


if __name__ == "__main__":
    main()
