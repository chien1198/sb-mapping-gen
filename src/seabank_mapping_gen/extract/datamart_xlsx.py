"""Extract dim/fact table definitions and transform logic from a datamart
design .xlsx file into the normalized DatamartDesign model.

NOTE: This is a skeleton. Real column positions/sheet layout need to be
confirmed against a sample file before `_parse_sheet` is finalized.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from seabank_mapping_gen.models import DatamartColumnLogic, DatamartDesign, DatamartTableDef

# Expected header names (case-insensitive), one sheet per dim/fact table.
EXPECTED_HEADERS = ["target column", "data type", "source table", "source column", "logic"]


def extract(xlsx_path: Path) -> DatamartDesign:
    workbook = load_workbook(str(xlsx_path), data_only=True)
    tables: list[DatamartTableDef] = []

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        table_def = _parse_sheet(sheet, sheet_name)
        if table_def is not None:
            tables.append(table_def)

    return DatamartDesign(source_file=str(xlsx_path), tables=tables)


def _parse_sheet(sheet, sheet_name: str) -> DatamartTableDef | None:
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < 2:
        return None

    header_row = [str(c).strip().lower() if c else "" for c in rows[0]]
    if not any(h in header_row for h in EXPECTED_HEADERS):
        return None

    columns: list[DatamartColumnLogic] = []
    for row in rows[1:]:
        if not any(row):
            continue
        columns.append(_row_to_column(header_row, row))

    table_type = "fact" if sheet_name.lower().startswith("fact") else "dim"
    return DatamartTableDef(table_name=sheet_name, table_type=table_type, columns=columns)


def _row_to_column(headers: list[str], row: tuple) -> DatamartColumnLogic:
    values = dict(zip(headers, row))

    def get(key: str) -> str:
        v = values.get(key)
        return str(v).strip() if v is not None else ""

    return DatamartColumnLogic(
        target_column=get("target column"),
        data_type=get("data type"),
        source_table=get("source table"),
        source_column=get("source column"),
        transform_logic=get("logic"),
    )
